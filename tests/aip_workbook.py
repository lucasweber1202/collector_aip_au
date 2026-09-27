"""Build small workbooks in the audited AIP TGP and retail layouts."""

from __future__ import annotations

import io
from datetime import date, datetime

import openpyxl

CITIES = ("Sydney", "Melbourne", "Brisbane", "Adelaide", "Perth", "Darwin", "Hobart", "National")
REGIONS = ("NSW", "VIC", "QLD", "SA", "WA", "NT", "TAS", "National\nAverage")


def tgp(days: list[date], base: float = 200.0) -> bytes:
    book = openpyxl.Workbook()
    first = book.active
    assert first is not None
    book.remove(first)
    for offset, name in enumerate(("Petrol TGP", "Diesel TGP")):
        sheet = book.create_sheet(name)
        sheet.append(["AVERAGE", *CITIES])
        for n, day in enumerate(days):
            sheet.append([datetime.combine(day, datetime.min.time()), *[base + offset * 30 + n + i / 10 for i in range(8)]])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def retail(years: list[int]) -> bytes:
    book = openpyxl.Workbook()
    first = book.active
    assert first is not None
    book.remove(first)
    for offset, name in enumerate(("Average Petrol Retail", "Average Diesel Retail")):
        sheet = book.create_sheet(name)
        sheet.append(["CALENDAR YEAR", *REGIONS])
        for n, year in enumerate(years):
            sheet.append([year, *[150.0 + offset * 10 + n + i for i in range(8)]])
        sheet.append(["2024-25", *[1.0] * 8])  # financial-year rows are a different measure
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
